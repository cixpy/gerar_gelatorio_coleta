import csv
import os
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, time, timedelta
from tkinter import Tk, filedialog, messagebox

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("Biblioteca openpyxl não encontrada. Instale com: pip install openpyxl")


def normalizar(texto):
    texto = "" if texto is None else str(texto)
    texto = unicodedata.normalize("NFD", texto.upper())
    return "".join(c for c in texto if unicodedata.category(c) != "Mn").strip()


def achar_coluna(cabecalhos, alternativas):
    mapa = {normalizar(h): h for h in cabecalhos if h is not None}
    for alt in alternativas:
        n = normalizar(alt)
        if n in mapa:
            return mapa[n]
    # busca parcial como fallback
    for hnorm, original in mapa.items():
        if any(normalizar(alt) in hnorm for alt in alternativas):
            return original
    return None


def segundos_de_tempo(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, timedelta):
        return valor.total_seconds()
    if isinstance(valor, time):
        return valor.hour * 3600 + valor.minute * 60 + valor.second + valor.microsecond / 1e6
    if isinstance(valor, datetime):
        return valor.hour * 3600 + valor.minute * 60 + valor.second + valor.microsecond / 1e6
    if isinstance(valor, (int, float)):
        # Excel armazena horário como fração de um dia.
        return float(valor) * 86400

    s = str(valor).strip().replace(",", ".")
    if not s:
        return None
    partes = s.split(":")
    try:
        if len(partes) == 3:
            h, m, seg = partes
            return float(h) * 3600 + float(m) * 60 + float(seg)
        if len(partes) == 2:
            m, seg = partes
            return float(m) * 60 + float(seg)
        return float(s)
    except ValueError:
        return None


def ler_csv(caminho):
    dados_bytes = open(caminho, "rb").read()
    texto = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            texto = dados_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    if texto is None:
        raise ValueError("Não foi possível identificar a codificação do CSV.")

    amostra = texto[:5000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
        delimitador = dialeto.delimiter
    except csv.Error:
        delimitador = ";"

    linhas = texto.splitlines()
    leitor = csv.DictReader(linhas, delimiter=delimitador)
    return list(leitor), leitor.fieldnames or []


def ler_excel(caminho):
    wb = load_workbook(caminho, data_only=True, read_only=True)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return [], []

    # Procura a linha que mais se parece com um cabeçalho do relatório.
    indice_header = 0
    for i, linha in enumerate(linhas[:30]):
        norm = [normalizar(v) for v in linha if v is not None]
        if any("FUNCION" in v for v in norm) and any("ATEND" in v or "FILA" in v for v in norm):
            indice_header = i
            break

    headers = [str(v).strip() if v is not None else "" for v in linhas[indice_header]]
    dados = []
    for linha in linhas[indice_header + 1:]:
        if not any(v not in (None, "") for v in linha):
            continue
        dados.append({headers[j]: linha[j] if j < len(linha) else None for j in range(len(headers)) if headers[j]})
    return dados, [h for h in headers if h]


def carregar(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".csv":
        return ler_csv(caminho)
    if ext in (".xlsx", ".xlsm"):
        return ler_excel(caminho)
    raise ValueError("Formato não suportado. Use CSV, XLSX ou XLSM.")


def extrair_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if valor is None:
        return None
    s = str(valor).strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.search(r"(\d{2}/\d{2}/\d{4})", s)
    if m:
        try:
            return datetime.strptime(m.group(1), "%d/%m/%Y").date()
        except ValueError:
            pass
    return None


def gerar_relatorio(arquivo_entrada, arquivo_saida):
    registros, headers = carregar(arquivo_entrada)
    if not registros:
        raise ValueError("O arquivo selecionado não possui registros.")

    col_fila = achar_coluna(headers, ["Fila", "Serviço", "Servico"])
    col_func = achar_coluna(headers, ["Funcionário", "Funcionario", "Atendente", "Operador"])
    col_tempo = achar_coluna(headers, ["Tempo de Atendimento", "Tempo Atendimento", "Duração Atendimento", "Duracao Atendimento"])
    col_nota = achar_coluna(headers, ["Nota", "Avaliação", "Avaliacao", "Satisfação", "Satisfacao"])
    col_data = achar_coluna(headers, ["Data", "Data Atendimento", "Início Atendimento", "Inicio Atendimento", "Data/Hora"])

    faltando = []
    if not col_fila: faltando.append("Fila")
    if not col_func: faltando.append("Funcionário")
    if not col_tempo: faltando.append("Tempo de Atendimento")
    if faltando:
        raise ValueError("Colunas obrigatórias não encontradas: " + ", ".join(faltando) + ".")

    resumo = defaultdict(lambda: {"tempos": [], "atendimentos": 0, "negativas": 0})
    datas = []

    for r in registros:
        fila = normalizar(r.get(col_fila))
        # Aceita tanto 'IIRGD - COLETA DE IMAGEM' quanto variações contendo COLETA DE IMAGEM.
        if "COLETA DE IMAGEM" not in fila:
            continue

        nome = str(r.get(col_func) or "").strip()
        if not nome:
            continue

        tempo = segundos_de_tempo(r.get(col_tempo))
        resumo[nome]["atendimentos"] += 1
        if tempo is not None:
            resumo[nome]["tempos"].append(tempo)

        if col_nota:
            nota = normalizar(r.get(col_nota))
            if nota in {"REGULAR", "RUIM", "NAO OPINOU"}:
                resumo[nome]["negativas"] += 1

        if col_data:
            d = extrair_data(r.get(col_data))
            if d:
                datas.append(d)

    if not resumo:
        raise ValueError("Nenhum registro da fila COLETA DE IMAGEM foi encontrado.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório COLETA"
    ws.sheet_view.showGridLines = False

    inicio = min(datas).strftime("%d/%m/%Y") if datas else ""
    fim = max(datas).strftime("%d/%m/%Y") if datas else ""
    periodo = f" - {inicio} a {fim}" if inicio and fim else ""

    ws.merge_cells("A1:D1")
    ws["A1"] = "Relatório COLETA" + periodo
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    cabecalhos = ["NOME", "MÉDIA/TEMPO DE ATENDIMENTO", "ATENDIMENTOS", "REGULAR/RUIM/NÃO OPINOU"]
    for c, titulo in enumerate(cabecalhos, 1):
        cell = ws.cell(row=2, column=c, value=titulo)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    nomes = sorted(resumo.keys(), key=lambda x: normalizar(x))
    linha = 3
    for nome in nomes:
        info = resumo[nome]
        media = sum(info["tempos"]) / len(info["tempos"]) if info["tempos"] else None
        ws.cell(linha, 1, nome.upper())
        if media is not None:
            # Grava como fração do dia para o Excel exibir duração corretamente.
            ws.cell(linha, 2, media / 86400)
            ws.cell(linha, 2).number_format = "[h]:mm:ss"
        ws.cell(linha, 3, info["atendimentos"])
        ws.cell(linha, 4, info["negativas"])
        linha += 1

    total_row = linha
    ws.cell(total_row, 1, "TOTAL DE ATENDIMENTOS")
    ws.cell(total_row, 3, sum(v["atendimentos"] for v in resumo.values()))
    ws.cell(total_row, 4, sum(v["negativas"] for v in resumo.values()))

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E7E6E6")

    for row in ws.iter_rows(min_row=2, max_row=total_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", horizontal="center" if cell.column > 1 else "left", wrap_text=True)

    for cell in ws[2]:
        cell.fill = header_fill
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = total_fill

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 32
    larguras = {1: 43, 2: 29, 3: 16, 4: 28}
    for col, largura in larguras.items():
        ws.column_dimensions[get_column_letter(col)].width = largura

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:D{total_row - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:D{total_row}"
    ws.oddFooter.center.text = "Página &P de &N"

    wb.save(arquivo_saida)
    return len(nomes), sum(v["atendimentos"] for v in resumo.values()), sum(v["negativas"] for v in resumo.values())


def main():
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    entrada = filedialog.askopenfilename(
        title="Selecione o relatório bruto de rastreamento",
        filetypes=[
            ("Arquivos suportados", "*.csv *.xlsx *.xlsm"),
            ("CSV", "*.csv"),
            ("Excel", "*.xlsx *.xlsm"),
            ("Todos os arquivos", "*.*"),
        ],
    )
    if not entrada:
        return

    pasta = os.path.dirname(entrada)
    nome_padrao = "Relatorio_COLETA.xlsx"
    saida = filedialog.asksaveasfilename(
        title="Salvar relatório COLETA",
        initialdir=pasta,
        initialfile=nome_padrao,
        defaultextension=".xlsx",
        filetypes=[("Planilha Excel", "*.xlsx")],
    )
    if not saida:
        return

    try:
        funcionarios, atendimentos, negativas = gerar_relatorio(entrada, saida)
        messagebox.showinfo(
            "Relatório concluído",
            f"Relatório gerado com sucesso!\n\n"
            f"Funcionários: {funcionarios}\n"
            f"Atendimentos: {atendimentos}\n"
            f"Regular/Ruim/Não opinou: {negativas}\n\n"
            f"Arquivo:\n{saida}",
        )
    except Exception as e:
        messagebox.showerror("Erro", str(e))


if __name__ == "__main__":
    main()
